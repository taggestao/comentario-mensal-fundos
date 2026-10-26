"""Gera a imagem da tabela de resultados do documento mensal a partir do Excel do Morningstar.

Uso: python3 tabela_resultados.py RetornosCarteirasFoFs_AAAAMM.xlsx saida.png

O Excel fica em G:\\Drives compartilhados\\Gestao_Base_Dados\\OFF\\Resultados Offshore
(no Drive: pasta "Resultados Offshore"). Se a equipe já tiver preparado uma aba limpa
("Sheet1 (2)"), ela é usada como está; senão, parte da aba exportada e aplica o mesmo
tratamento: tira as colunas 1M (-1) e 1M (-2) e o grupo ALT Onshore.
Cores: escala de 3 cores por coluna de retorno, igual à formatação condicional do Excel
(mínimo F8696B, percentil 50 FCFCFF, máximo 63BE7B).
"""
import datetime
import sys

import openpyxl
from PIL import Image, ImageDraw, ImageFont

GRUPOS_FORA = {'ALT Onshore'}
COLUNAS_FORA = {'1M (-1)', '1M (-2)'}
FONTES = '/usr/share/fonts/truetype/liberation/'
ESCALA = 3  # pixels por ponto (≈ 216 dpi)
CORES = ((0xF8, 0x69, 0x6B), (0xFC, 0xFC, 0xFF), (0x63, 0xBE, 0x7B))


def ler(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    limpa = len(wb.worksheets) > 1
    ws = wb.worksheets[-1] if limpa else wb.worksheets[0]
    linhas = [[c.value for c in r] for r in ws.iter_rows()]
    i_per = next(i for i, l in enumerate(linhas) if '1M' in l)
    i_cab = next(i for i, l in enumerate(linhas) if l and l[0] == 'Group/Investment')
    periodos, cab = linhas[i_per], linhas[i_cab]
    manter = [j for j in range(len(cab)) if cab[j] is not None and periodos[j] not in COLUNAS_FORA]
    corpo, pular = [], False
    for l in linhas[i_cab + 1:]:
        nome = l[0]
        if nome is not None and all(v is None for v in l[1:]) and nome != 'Number of investments ranked':
            pular = nome in GRUPOS_FORA          # título de grupo
        if not pular:
            corpo.append([l[j] if j < len(l) else None for j in manter])
    while corpo and all(v is None for v in corpo[-1]):
        corpo.pop()
    return [periodos[j] for j in manter], [cab[j] for j in manter], corpo


def cor(v, vmin, vmed, vmax):
    if v <= vmed:
        a, b, t = CORES[0], CORES[1], 0 if vmed == vmin else (v - vmin) / (vmed - vmin)
    else:
        a, b, t = CORES[1], CORES[2], 0 if vmax == vmed else (v - vmed) / (vmax - vmed)
    return tuple(round(x + (y - x) * t) for x, y in zip(a, b))


def mediana(xs):
    xs = sorted(xs)
    n = len(xs)
    # percentil 50 do Excel (interpolação linear)
    p = (n - 1) * 0.5
    k = int(p)
    return xs[k] + (xs[min(k + 1, n - 1)] - xs[k]) * (p - k)


def fmt(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return '' if v is None else str(v)


def desenhar(periodos, cab, corpo, saida):
    pt = lambda x: round(x * ESCALA)
    f = {k: ImageFont.truetype(FONTES + arq, pt(6.5)) for k, arq in
         (('n', 'LiberationSans-Regular.ttf'), ('b', 'LiberationSans-Bold.ttf'), ('i', 'LiberationSans-Italic.ttf'))}
    num = [j for j, p in enumerate(periodos) if p]
    largura = lambda t, fonte: f[fonte].getlength(t)
    # largura de cada coluna pelo maior texto
    ws = []
    for j in range(len(cab)):
        textos = [fmt(l[j]) for l in corpo if not (j == 0 and l[1] is None)] + cab[j].split('\n') + [fmt(periodos[j])]
        ws.append(max(largura(t, 'n') for t in textos) + pt(8))
    ws[0] = max(ws[0], max(largura(fmt(l[0]), 'b') for l in corpo) + pt(8))
    alt, alt_cab, margem = pt(10), pt(30), pt(6)
    W = int(sum(ws) + 2 * margem)
    H = int(margem * 2 + alt + alt_cab + alt * len(corpo))
    img = Image.new('RGB', (W, H), 'white')
    d = ImageDraw.Draw(img)
    xs = [margem]
    for w in ws:
        xs.append(xs[-1] + w)

    # escalas de cor por coluna de retorno (todas as linhas numéricas, como no Excel)
    faixas = {}
    for j in num:
        vals = [l[j] for l in corpo if isinstance(l[j], (int, float))]
        if vals:
            faixas[j] = (min(vals), mediana(vals), max(vals))

    def texto(j, y, t, fonte='n', h=alt, direita=None):
        direita = j in num or j == len(cab) - 1 or j == 6 if direita is None else direita
        tw = largura(t, fonte)
        x = xs[j + 1] - pt(3) - tw if direita else xs[j] + pt(3)
        d.text((x, y + (h - pt(6.5)) / 2 - pt(1)), t, font=f[fonte], fill='black')

    y = margem
    for j in num:
        texto(j, y, periodos[j], 'b', direita=False)
    y += alt
    for j, c in enumerate(cab):
        partes = c.split('\n')
        for k, p in enumerate(partes):
            yy = y + alt_cab - alt * (len(partes) - k)
            texto(j, yy, p, direita=j in num)
    y += alt_cab
    d.line((xs[0], y, xs[-1], y), fill='black', width=max(1, pt(0.5)))
    for l in corpo:
        nome = l[0]
        if nome is not None and all(v is None for v in l[1:]) and nome != 'Number of investments ranked':
            texto(0, y, nome, 'b', direita=False)
        elif nome == 'Number of investments ranked' or (nome or '').startswith('Benchmark'):
            for j, v in enumerate(l):
                if j in faixas and isinstance(v, (int, float)):
                    d.rectangle((xs[j], y, xs[j + 1] - 1, y + alt - 1), fill=cor(v, *faixas[j]))
                if v is not None:
                    texto(j, y, fmt(v), 'i' if j == 0 else 'n')
        else:
            for j, v in enumerate(l):
                if j in faixas and isinstance(v, (int, float)):
                    d.rectangle((xs[j], y, xs[j + 1] - 1, y + alt - 1), fill=cor(v, *faixas[j]))
                if v is not None:
                    texto(j, y, fmt(v))
        y += alt
    # paleta de 256 cores: o arquivo fica bem menor sem perda visível
    img.convert('P', palette=Image.ADAPTIVE, colors=256).save(saida, dpi=(72 * ESCALA, 72 * ESCALA), optimize=True)
    return img.size


if __name__ == '__main__':
    periodos, cab, corpo = ler(sys.argv[1])
    datas = {l[6] for l in corpo if isinstance(l[6], datetime.datetime)}
    print('linhas:', len(corpo), '| data de retorno:', ', '.join(sorted(fmt(x) for x in datas)))
    print('tamanho:', desenhar(periodos, cab, corpo, sys.argv[2]))
